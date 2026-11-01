"""Build doc/_build/reference.docx - the style template every .docx in doc/ uses.

Run:  python doc/_build/make_reference.py
Needs: pypandoc-binary (bundled pandoc) and python-docx.
"""

import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
BUILD = ROOT / "doc" / "_build"

INK = RGBColor(0x0F, 0x17, 0x2A)      # slate-900 - headings
INK_SOFT = RGBColor(0x33, 0x41, 0x55)  # slate-700
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)    # blue-700 - H2, links
MUTED = RGBColor(0x64, 0x74, 0x8B)     # slate-500 - subtitle, caption
BORDER = "94A3B8"                       # slate-400 - table rules
HEADER_FILL = "EEF2F7"                  # table header shading


def dump_default_reference() -> Path:
    target = BUILD / "reference-default.docx"
    pandoc = pypandoc.get_pandoc_path()
    with open(target, "wb") as handle:
        subprocess.run(
            [pandoc, "--print-default-data-file", "reference.docx"],
            check=True, stdout=handle,
        )
    return target


def patch_table_style(docx_path: Path) -> None:
    """The stock reference has a borderless 'Table' style; give tables real rules."""
    with zipfile.ZipFile(docx_path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
        styles = entries["word/styles.xml"].decode("utf-8")

    borders = (
        '<w:tblBorders>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{BORDER}"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{BORDER}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{BORDER}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{BORDER}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{BORDER}"/>'
        '</w:tblBorders>'
        '<w:tblCellMar>'
        '<w:top w:w="60" w:type="dxa"/><w:left w:w="90" w:type="dxa"/>'
        '<w:bottom w:w="60" w:type="dxa"/><w:right w:w="90" w:type="dxa"/>'
        '</w:tblCellMar>'
    )
    # 'Table' is the styleId pandoc gives every pipe table.
    styles, hits = re.subn(
        r'(<w:style w:type="table"[^>]*w:styleId="Table">)(.*?)(</w:style>)',
        lambda m: m.group(1) + re.sub(r"<w:tblPr>.*?</w:tblPr>", "", m.group(2), flags=re.S)
        + f"<w:tblPr>{borders}</w:tblPr>" + m.group(3),
        styles,
        flags=re.S,
    )
    if hits != 1:
        raise SystemExit(f"expected exactly one 'Table' style, patched {hits}")

    # pandoc's table style already carries a firstRow conditional; put the header
    # fill there so Word shades the header row of every pipe table.
    styles, header_hits = re.subn(
        r'(<w:tblStylePr w:type="firstRow">\s*<w:tcPr>)',
        r'\1<w:shd w:val="clear" w:color="auto" w:fill="' + HEADER_FILL + '"/>',
        styles,
        count=1,
    )

    entries["word/styles.xml"] = styles.encode("utf-8")
    print(f"  table rules patched: {hits}, header shading patched: {header_hits}")
    with zipfile.ZipFile(docx_path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, blob in entries.items():
            out.writestr(name, blob)


def main() -> None:
    source = dump_default_reference()
    target = BUILD / "reference.docx"
    document = Document(str(source))

    for section in document.sections:
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.9)
        section.left_margin = Cm(2.1)
        section.right_margin = Cm(2.1)

    styles = document.styles
    by_name = {style.name: style for style in styles}
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.12

    heading_sizes = {"Heading 1": 15, "Heading 2": 12.5, "Heading 3": 11.5, "Heading 4": 11}
    heading_colors = {
        "Heading 1": INK,
        "Heading 2": ACCENT,
        "Heading 3": INK_SOFT,
        "Heading 4": INK_SOFT,
    }
    for name, size in heading_sizes.items():
        style = styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = heading_colors[name]
        style.font.name = "Calibri"
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    styles["Title"].font.size = Pt(26)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = INK
    styles["Subtitle"].font.size = Pt(11.5)
    styles["Subtitle"].font.color.rgb = MUTED
    styles["Author"].font.size = Pt(10.5)
    styles["Author"].font.color.rgb = MUTED
    styles["Date"].font.size = Pt(10.5)
    styles["Date"].font.color.rgb = MUTED

    for name in ("Source Code", "Verbatim Char"):
        if name in by_name:
            by_name[name].font.name = "Consolas"
            by_name[name].font.size = Pt(9)
    if "Caption" in by_name:
        by_name["Caption"].font.size = Pt(9)
        by_name["Caption"].font.color.rgb = MUTED
    for bullet in ("Compact", "First Paragraph", "Body Text"):
        if bullet in by_name:
            by_name[bullet].font.size = Pt(10.5)

    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Asset Management App  |  documentation set  |  page ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    field = paragraph.add_run()
    field.font.size = Pt(8)
    field.font.color.rgb = MUTED
    simple = paragraph._p.makeelement(qn("w:fldSimple"), {qn("w:instr"): " PAGE "})
    paragraph._p.append(simple)

    document.save(str(target))
    patch_table_style(target)

    check = Document(str(target))
    print("reference.docx written:", target)
    print("  margins cm:", round(check.sections[0].left_margin.cm, 2))
    print("  Normal pt:", check.styles["Normal"].font.size.pt)
    print("  Heading 1 pt/color:", check.styles["Heading 1"].font.size.pt,
          check.styles["Heading 1"].font.color.rgb)
    print("  footer text:", check.sections[0].footer.paragraphs[0].text)


if __name__ == "__main__":
    sys.exit(main())
