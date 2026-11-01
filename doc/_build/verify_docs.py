"""Verify the rendered .docx set: structure, footer, and a secret scan."""

import re
import sys
import zipfile
from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parents[1]
SECRET_PATTERNS = {
    "email-looking": re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"),
    "password-value": re.compile(r"(?i)(password|secret|token)\s*[:=]\s*[^\s<]{8,}"),
    "private-key": re.compile(r"BEGIN (RSA |EC )?PRIVATE KEY"),
    "aws-ish": re.compile(r"(?i)AKIA[0-9A-Z]{16}"),
    "ipv4-ish": re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b"),
}
ALLOWED = {
    "127.0.0.1",
    "192.168.1.23",
    "192.168.1.0",
    "192.168.7.0",
    "10.0.0.5",
    "10.0.0.40",
    "172.17.0.1",
    "192.168.1.10",
    "0.0.0.0",
    "1.1.1.1",
    "smtp.example.com",
    "you@example.com",
    "a@x.com",
    "b@x.com",
    "192.168.1.20",
    "token = crypto.randomBytes(32).toString('hex')",
}


def text_of(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def main():
    bad = []
    for docx in sorted(OUT.glob("*.docx")):
        doc = Document(str(docx))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        body = text_of(doc)
        with zipfile.ZipFile(docx) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
            footers = [n for n in archive.namelist() if n.startswith("word/footer")]
            footer_xml = "".join(archive.read(n).decode("utf-8") for n in footers)
        page_field = "PAGE" in footer_xml
        toc = 'TOC' in xml
        # a markdown pipe-table row that failed to render
        stray = re.findall(r"^\s*\|.*\|\s*$", body, re.M)
        stray_bold = re.findall(r"\*\*[^*\n]{2,}\*\*", body)
        stray_code = re.findall(r"(?<![`])`[^`\n]{2,}`(?!`)", body)
        findings = []
        for name, pattern in SECRET_PATTERNS.items():
            hits = {h.group(0) for h in pattern.finditer(body)} - ALLOWED
            if hits:
                findings.append(f"{name}: {sorted(hits)[:6]}")
        print(
            f"{docx.name:42s} paras={len(doc.paragraphs):4d} tables={len(doc.tables):3d} "
            f"H1={len(h1):3d} H2={len(headings) - len(h1):3d} chars={len(body):6,d} "
            f"toc={toc} page={page_field} stray_table={len(stray)} stray_bold={len(stray_bold)} "
            f"stray_code={len(stray_code)}"
        )
        if stray or stray_bold or stray_code:
            bad.append(f"{docx.name}: unrendered markdown -> {stray[:2]} {stray_bold[:2]} {stray_code[:3]}")
        if not toc or not page_field:
            bad.append(f"{docx.name}: missing toc={toc} page-field={page_field}")
        if findings:
            bad.append(f"{docx.name}: possible secrets -> {findings}")
    print("\nproblems:" if bad else "\nno problems found")
    for line in bad:
        print(" -", line)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
